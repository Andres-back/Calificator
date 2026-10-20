"""Deterministic evaluation material for documents and vision consumers.

Only private, revision-scoped bytes enter this module. No URL fetching, HTML
injection or inference is used to render numerical data.
"""
from __future__ import annotations

import base64
from dataclasses import dataclass
from html import escape
from io import BytesIO
from math import cos, pi, sin
from textwrap import wrap
from uuid import UUID

from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import select

from app.modules.evaluaciones.models import EvaluacionRecurso
from app.modules.evaluaciones.visual_schemas import VisualBlock, VisualChart, VisualManifest, VisualTable
from app.modules.evaluaciones.visual_service import inspect_resource_bytes, read_resource_bytes

PALETTE = ['#4338ca', '#0f766e', '#be123c', '#a16207', '#0369a1', '#7e22ce']


def _lines(value: str, width: int) -> list[str]:
    lines = [line for paragraph in value.splitlines() or ['']
             for line in wrap(paragraph, width=width, break_long_words=True) or ['']]
    if len(lines) > 10:
        raise ValueError('El texto es demasiado extenso para esta figura. Divide la tabla o utiliza un texto de apoyo.')
    return lines


def render_structured_visual(value: VisualTable | VisualChart) -> bytes:
    """Lossless PNG derivative, including values, labels and units; bounded work."""
    font = ImageFont.load_default(size=20)
    small = ImageFont.load_default(size=32)
    width = 1500
    if isinstance(value, VisualTable):
        if len(value.encabezados) > 10 or len(value.filas) > 60:
            raise ValueError('La tabla es demasiado amplia para una figura legible. Divide sus datos en tablas más pequeñas.')
        columns = len(value.encabezados)
        cell_width = width // columns
        rows = [value.encabezados, *value.filas]
        cells = [[_lines(cell, max(8, cell_width // 12)) for cell in row] for row in rows]
        heights = [max(len(cell) for cell in row) * 27 + 26 for row in cells]
        height = sum(heights) + (48 if value.unidad else 0)
        if height > 9000:
            raise ValueError('La tabla necesita más espacio para leerse. Divide sus datos antes de continuar.')
        image = Image.new('RGB', (width, height), 'white')
        draw = ImageDraw.Draw(image)
        top = 0
        for index, row in enumerate(cells):
            bottom = top + heights[index]
            for column, lines in enumerate(row):
                left = column * cell_width
                right = width - 1 if column == columns - 1 else left + cell_width
                draw.rectangle((left, top, right, bottom - 1), fill='#e0e7ff' if index == 0 else 'white', outline='#94a3b8', width=2)
                draw.multiline_text((left + 14, top + 12), '\n'.join(lines), font=font, fill='#0f172a', spacing=5)
            top = bottom
        if value.unidad:
            draw.text((14, top + 12), f'Unidad: {value.unidad}', font=font, fill='#0f172a')
    else:
        if len(value.categorias) > 30 or len(value.series) > len(PALETTE):
            raise ValueError('La gráfica necesita demasiadas etiquetas. Divide sus datos para que pueda leerse.')
        if any(len(text) > 90 for text in [*value.categorias, *(series.nombre for series in value.series), value.unidad or '', value.eje_x or '', value.eje_y or '']):
            raise ValueError('Las etiquetas son demasiado extensas. Divide o simplifica el texto de la gráfica.')
        legend_labels = [_lines(f'{index + 1}. {category}: {number} ({number / sum(value.series[0].valores):.1%})', 28)
                         for index, (category, number) in enumerate(zip(value.categorias, value.series[0].valores))] if value.tipo == 'circular' else []
        legend_height = sum(len(lines) * 40 + 18 for lines in legend_labels)
        height = max(1120, 180 + legend_height)
        image = Image.new('RGB', (width, height), 'white')
        draw = ImageDraw.Draw(image)
        if value.tipo == 'circular':
            total = sum(value.series[0].valores)
            angle = -90.0
            bounds = (70, 140, 830, 900)
            legend_top = 135
            for index, (category, number) in enumerate(zip(value.categorias, value.series[0].valores)):
                end = angle + number / total * 360
                color = PALETTE[index % len(PALETTE)]
                draw.pieslice(bounds, start=angle, end=end, fill=color, outline='white', width=3)
                draw.rectangle((880, legend_top, 900, legend_top + 20), fill=color)
                draw.multiline_text((915, legend_top), '\n'.join(legend_labels[index]), font=small, fill='#0f172a', spacing=5)
                legend_top += len(legend_labels[index]) * 40 + 18
                if number > 0:
                    middle = (angle + end) / 2 * pi / 180
                    draw.text((440 + cos(middle) * 240, 510 + sin(middle) * 240), str(index + 1), font=font, fill='white')
                angle = end
        else:
            left, top, right, bottom = 120, 150, 1430, 710
            values = [number for series in value.series for number in series.valores]
            low, high = min(0, min(values)), max(0, max(values))
            scale = max(abs(low), abs(high)) or 1
            span = high / scale - low / scale or 1
            def y(number):
                return bottom - (number / scale - low / scale) / span * (bottom - top)
            for tick in range(6):
                number = (low / scale + span * tick / 5) * scale
                position = y(number)
                draw.line((left, position, right, position), fill='#cbd5e1', width=1)
                draw.text((12, position - 10), f'{number:.4g}', font=small, fill='#334155')
            draw.line((left, y(0), right, y(0)), fill='#334155', width=2)
            step = (right - left) / len(value.categorias)
            for series_index, series in enumerate(value.series):
                color = PALETTE[series_index]
                points = []
                for index, number in enumerate(series.valores):
                    center = left + (index + 0.5) * step
                    if value.tipo == 'barras':
                        bar_width = step * 0.75 / len(value.series)
                        x = center - step * 0.375 + series_index * bar_width
                        draw.rectangle((x, min(y(number), y(0)), x + bar_width, max(y(number), y(0))), fill=color)
                    else:
                        points.append((center, y(number)))
                        draw.ellipse((center - 5, y(number) - 5, center + 5, y(number) + 5), fill=color)
                if len(points) > 1:
                    draw.line(points, fill=color, width=4)
                draw.text((left, 900 + series_index * 30), series.nombre, font=font, fill=color)
            for index, category in enumerate(value.categorias):
                # Labels and exact values are also available in the data table.
                label = f'{index + 1}' if len(value.categorias) > 10 or len(category) > 24 else '\n'.join(_lines(category, 12))
                draw.multiline_text((left + (index + 0.1) * step, bottom + 15), label, font=small, fill='#0f172a', spacing=4)
            # Full category names and exact values remain in the raster too.
            data_image = Image.open(BytesIO(render_structured_visual(VisualTable(
                encabezados=[value.eje_x or 'Categoría', *(series.nombre for series in value.series)],
                filas=[[category, *(str(series.valores[index]) for series in value.series)]
                       for index, category in enumerate(value.categorias)], unidad=value.unidad,
            ))))
            combined = Image.new('RGB', (width, height + data_image.height), 'white')
            combined.paste(image, (0, 0))
            combined.paste(data_image, (0, height))
            image = combined
            draw = ImageDraw.Draw(image)
        heading = ' · '.join(text for text in [value.eje_y, value.unidad] if text)
        draw.text((70, 30), heading, font=font, fill='#0f172a')
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


@dataclass(frozen=True)
class ResolvedVisualBlock:
    block: VisualBlock
    image_bytes: bytes | None = None

    def __post_init__(self):
        if self.block.seguridad_contenido != 'publicable_confirmado' or not self.block.legible or not self.block.pregunta_ids:
            raise ValueError('Revisa el material y confirma que no contiene respuestas del estudiante.')
        if self.block.tipo != 'texto':
            if not self.image_bytes:
                raise ValueError('Falta material visual de la pregunta. Vuelve a cargarlo antes de descargar.')
            inspect_resource_bytes(self.image_bytes, clase='imagen')


async def resolve_visual_material(db, evaluation_id: UUID, revision) -> list[ResolvedVisualBlock]:
    """Caller has already authorized the requested evaluation/revision."""
    raw = revision.snapshot_json.get('contenido_visual')
    if raw is None:
        return []
    manifest = VisualManifest.model_validate(raw)
    if manifest.cobertura.estado != 'completa':
        raise ValueError('Revisa las figuras y las preguntas a las que pertenecen antes de descargar la evaluación completa.')
    blocks = [block for block in manifest.bloques if block.seguridad_contenido != 'privado_original']
    required = {block.recurso_id for block in blocks if block.recurso_id is not None}
    resources = {resource.id: resource for resource in await db.scalars(select(EvaluacionRecurso).where(
        EvaluacionRecurso.id.in_(required), EvaluacionRecurso.evaluacion_id == evaluation_id,
        EvaluacionRecurso.estado == 'ready',
    ))} if required else {}
    resolved = []
    cache = {}
    for block in sorted(blocks, key=lambda item: item.orden):
        data = None
        if block.recurso_id is not None:
            resource = resources.get(block.recurso_id)
            if resource is None or str(resource.id) not in revision.recurso_ids or resource.evaluacion_id != evaluation_id or resource.estado != 'ready' or resource.clase in {'original', 'pagina'} or resource.seguridad_contenido != 'publicable_confirmado':
                raise ValueError('Falta material confirmado de esta evaluación. Revisa sus imágenes antes de descargar.')
            if resource.id not in cache:
                cache[resource.id] = await read_resource_bytes(resource)
            data = cache[resource.id]
            if block.tabla or block.grafica:
                canonical = render_structured_visual(block.tabla or block.grafica)
                if canonical != data:
                    raise ValueError('Los datos y la figura no coinciden. Revisa el material antes de descargar.')
        resolved.append(ResolvedVisualBlock(block, data))
    return resolved


def visual_block_html(item: ResolvedVisualBlock) -> str:
    block = item.block
    description = escape(block.descripcion)
    if block.tipo == 'texto':
        body = '<p style="white-space:pre-wrap">' + escape(block.texto or '') + '</p>'
    elif block.tabla:
        body = '<table class="grid"><thead><tr>' + ''.join('<th>' + escape(cell) + '</th>' for cell in block.tabla.encabezados) + '</tr></thead><tbody>'
        body += ''.join('<tr>' + ''.join('<td>' + escape(cell) + '</td>' for cell in row) + '</tr>' for row in block.tabla.filas) + '</tbody></table>'
        if block.tabla.unidad:
            body += '<p>Unidad: ' + escape(block.tabla.unidad) + '</p>'
    else:
        metadata = inspect_resource_bytes(item.image_bytes, clase='imagen')
        encoded = base64.b64encode(item.image_bytes).decode('ascii')
        body = f'<img src="data:{metadata["mime"]};base64,{encoded}" alt="{description}" style="display:block;max-width:100%;max-height:320px;width:440px;height:auto;object-fit:contain;margin:auto"/>'
        if block.grafica:
            body += '<table class="grid"><thead><tr><th>Categoría</th>' + ''.join('<th>' + escape(series.nombre) + '</th>' for series in block.grafica.series) + '</tr></thead><tbody>'
            for index, category in enumerate(block.grafica.categorias):
                body += '<tr><td>' + escape(category) + '</td>' + ''.join(f'<td>{series.valores[index]}</td>' for series in block.grafica.series) + '</tr>'
            body += '</tbody></table>'
    caption = escape(block.pie or block.descripcion)
    return '<figure style="margin:12px 0;break-inside:avoid">' + body + (f'<figcaption>{caption}</figcaption>' if caption else '') + '</figure>'


def append_visual_docx(document, item: ResolvedVisualBlock) -> None:
    from docx.oxml import OxmlElement
    from docx.shared import Cm

    def style_table(table):
        table.style = 'Table Grid'
        borders = OxmlElement('w:tblBorders')
        for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + name)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                border.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}' + key, value)
            borders.append(border)
        table._tbl.tblPr.append(borders)

    block = item.block
    if block.tipo == 'texto':
        document.add_paragraph(block.texto)
    elif block.tabla:
        table = document.add_table(rows=1, cols=len(block.tabla.encabezados))
        style_table(table)
        for cell, text in zip(table.rows[0].cells, block.tabla.encabezados):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in block.tabla.filas:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = text
        if block.tabla.unidad:
            document.add_paragraph('Unidad: ' + block.tabla.unidad)
    else:
        metadata = inspect_resource_bytes(item.image_bytes, clase='imagen')
        width = min(12, 10 * metadata['ancho'] / metadata['alto'])
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = bool(block.pie or block.descripcion)
        picture = paragraph.add_run().add_picture(BytesIO(item.image_bytes), width=Cm(width))
        picture._inline.docPr.set('descr', block.descripcion or 'Material de la pregunta')
        if block.grafica:
            table = document.add_table(rows=1, cols=1 + len(block.grafica.series))
            style_table(table)
            for cell, text in zip(table.rows[0].cells, ['Categoría', *(series.nombre for series in block.grafica.series)]):
                cell.text = text
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for index, category in enumerate(block.grafica.categorias):
                for cell, text in zip(table.add_row().cells, [category, *(str(series.valores[index]) for series in block.grafica.series)]):
                    cell.text = text
    if block.pie or block.descripcion:
        paragraph = document.add_paragraph(block.pie or block.descripcion)
        paragraph.paragraph_format.keep_with_next = True
