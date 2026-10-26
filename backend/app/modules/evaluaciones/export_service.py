"""Documentos de solo lectura a partir del contenido guardado; nunca usa IA."""
from __future__ import annotations

from copy import deepcopy
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata
from urllib.parse import quote

from app.modules.evaluaciones.service import sanitize_student_payload
from app.modules.evaluaciones.visual_render import ResolvedVisualBlock, append_visual_docx
from app.modules.evaluaciones.reading_service import evaluation_reading
from app.modules.evaluaciones.print_layout import option_columns

MISSING_ANSWER = 'Sin respuesta esperada guardada'
WORD_TYPES = {'examen', 'quiz_rapido', 'taller', 'ficha', 'lectura_comprensiva'}


def _text(value: object) -> str:
    if isinstance(value, dict):
        return str(value.get('texto') or value.get('enunciado') or value.get('valor') or '')
    if isinstance(value, list):
        return '; '.join(_text(item) for item in value)
    return str(value) if value is not None else ''


def _answer(item: dict) -> object:
    return next((item[key] for key in ('respuesta', 'texto', 'respuesta_esperada', 'valor', 'respuesta_correcta')
                 if item.get(key) not in (None, '', [])), None)


def build_evaluation_material(evaluation, *, soluciones: bool, visual_blocks: list[ResolvedVisualBlock] | None = None) -> dict:
    from app.modules.evaluaciones.document_schemas import EvaluationDocument
    from app.modules.evaluaciones.document_structure import public_document

    document = getattr(evaluation, 'estructura_documento', None)
    projection = public_document(EvaluationDocument.model_validate(document), evaluation.preguntas,
        material_ids=[item.block.id for item in visual_blocks or []]) if document is not None else None
    expected = {str(answer.get('numero', index)): _answer(answer)
                for index, answer in enumerate(evaluation.respuestas_esperadas or [], 1)
                if isinstance(answer, dict)}
    questions = []
    pending_visuals = {str(item.block.id): item for item in visual_blocks or []}
    document_visuals = []
    section_totals = {}
    if projection:
        placed = {node['bloque_id'] for node in projection['nodos'] if node['tipo'] == 'material'}
        document_visuals = [pending_visuals.pop(identifier) for identifier in placed]
        if document.get('puntuacion_secciones'):
            by_node = {node['id']: node for node in projection['nodos']}
            by_question = {str(question['contenido_id']): Decimal(str(question['puntaje'])) for question in evaluation.preguntas}
            for node in projection['nodos']:
                if node['tipo'] == 'pregunta':
                    parent = node.get('padre_id')
                    while parent:
                        section_totals[parent] = section_totals.get(parent, Decimal('0')) + by_question[node['pregunta_id']]
                        parent = by_node[parent].get('padre_id')
    for index, raw in enumerate(evaluation.preguntas or [], 1):
        raw = raw if isinstance(raw, dict) else {'enunciado': str(raw)}
        clean = sanitize_student_payload(deepcopy(raw))
        clean = {key: value for key, value in clean.items() if not key.startswith('_print_')}
        number = raw.get('numero', index)
        clean['numero'] = number
        clean['opciones'] = [_text(option) for option in clean.get('opciones') or []]
        # Only trusted, internally resolved values enter the rendering channel.
        clean.pop('_evaluation_visuals', None)
        related = [item for item in pending_visuals.values()
                   if str(raw.get('contenido_id')) in {str(value) for value in item.block.pregunta_ids}]
        if related:
            clean['_evaluation_visuals'] = related
            for item in related:
                pending_visuals.pop(str(item.block.id))
        if soluciones:
            # El contrato canónico separado es la fuente de verdad de la clave.
            answer = expected.get(str(number))
            if answer is None:
                # Compatibilidad con claves antiguas guardadas dentro de la pregunta.
                answer = raw.get('respuesta_correcta')
            clean['respuesta_correcta'] = _text(answer) if answer is not None else MISSING_ANSWER
        questions.append(clean)
    if not questions:
        raise ValueError('La evaluación aún no tiene preguntas para visualizar o descargar.')
    if pending_visuals:
        raise ValueError('Revisa a qué preguntas pertenece el material antes de descargar la evaluación.')
    return {'tipo': 'examen', 'titulo': evaluation.nombre, 'created_at': evaluation.created_at,
            'contenido_json': {'titulo': evaluation.nombre,
                              **({'configuracion_impresion': evaluation.configuracion_impresion}
                                 if getattr(evaluation, 'configuracion_impresion', None) is not None else {}),
                              **({'documento': projection, '_document_visuals': document_visuals,
                                  'subtotales_secciones': {key: format(value, '.2f') for key, value in section_totals.items()}} if projection else {}),
                              'texto': None if projection else evaluation_reading(evaluation), 'impresion_compacta': True,
                              'instrucciones': evaluation.descripcion or 'Lee y responde cada punto.',
                              'preguntas': questions, 'total_puntaje': evaluation.nota_maxima}}


def document_disposition(name: str, extension: str, attachment: bool) -> str:
    name = ''.join(c for c in name if c.isalnum() or c in ' -_').strip()[:100] or 'evaluacion'
    ascii_name = unicodedata.normalize('NFKD', name).encode('ascii', 'ignore').decode().strip() or 'evaluacion'
    mode = 'attachment' if attachment else 'inline'
    return f'{mode}; filename="{ascii_name}.{extension}"; filename*=UTF-8\'\'{quote(name + "." + extension)}'


def _keep_short_word_question(document, start: int, question: dict) -> None:
    """Keep bounded questions/figures with their choices, not across page-sized text."""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    elements = list(document.element.body)[start:]
    elements = [element for element in elements if element.tag.endswith(('}p', '}tbl'))]
    if not elements or len(_text(question.get('enunciado'))) + sum(len(_text(option)) for option in question.get('opciones') or []) >= 650:
        return
    # Large figures/tables must be allowed to continue on another page.
    if len(question.get('_evaluation_visuals') or []) > 1 or any(
        len(table.xpath('./w:tr')) > 8 or any(len(''.join(cell.itertext())) > 180 for cell in table.xpath('.//w:tc'))
        for element in elements for table in ([element] if element.tag.endswith('}tbl') else [])
    ):
        return
    paragraphs = [p for element in elements for p in ([element] if element.tag.endswith('}p') else element.xpath('.//w:p'))]
    for element in paragraphs:
        paragraph = Paragraph(element, document)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
    for element in elements:
        if element.tag.endswith('}tbl'):
            for row in element.xpath('./w:tr'):
                row.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    last = elements[-1]
    ending = [last] if last.tag.endswith('}p') else last.xpath('./w:tr[last()]//w:p')
    for element in ending:
        Paragraph(element, document).paragraph_format.keep_with_next = False


def render_evaluation_docx(material: dict, *, soluciones: bool) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from app.modules.evaluaciones.print_layout import prepare_print_material
    from app.modules.evaluaciones.print_settings import resolve_print_settings
    from app.modules.evaluaciones.document_print import document_print_events

    material = prepare_print_material(material)

    kind = material.get('tipo')
    if kind not in WORD_TYPES:
        raise ValueError('Este material conserva su formato en PDF. Su conversión completa a Word aún no está disponible.')
    content = deepcopy(material.get('contenido_json') or {})
    question_key = {'taller': 'puntos', 'ficha': 'ejercicios'}.get(kind, 'preguntas')
    questions = content.get(question_key) or []
    if not questions:
        raise ValueError('El contenido aún no tiene preguntas para descargar en Word. Puedes consultar su PDF.')
    # No se omiten imágenes/contexto gráficos silenciosamente para aparentar una exportación íntegra.
    graphic_keys = ('imagen', 'imagen_url', 'imagenes', 'grafico', 'diagrama', 'ilustracion')
    if any(content.get(k) for k in graphic_keys) or any(isinstance(q, dict) and any(q.get(k) for k in graphic_keys) for q in questions):
        raise ValueError('La evaluación incluye contenido gráfico que debes conservar en PDF; Word no está disponible para este material.')
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[style_name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Word's theme can override an explicit font and add a blue title rule.
        # These printouts use a plain, portable typographic hierarchy.
        fonts = style.element.get_or_add_rPr().rFonts
        if fonts is not None:
            for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                fonts.attrib.pop(qn('w:' + key), None)
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):
                style.element.pPr.remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    print_settings = resolve_print_settings(content.get('configuracion_impresion'))
    compact = bool(content.get('impresion_compacta'))
    if compact:
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.05
    if print_settings:
        normal.font.size = Pt(print_settings.fuente_pt)
        normal.paragraph_format.space_after = Pt({'compacta': 2, 'media': 5, 'amplia': 8}[print_settings.distribucion])
        document.styles['Title'].font.size = Pt(print_settings.fuente_pt + 5.5)
        document.styles['Title'].paragraph_format.space_after = Pt(3)
        if not print_settings.blanco_negro:
            document.styles['Title'].font.color.rgb = RGBColor.from_string('4338CA')
        section.page_width, section.page_height = Cm(21.59), Cm(27.94)
        section.left_margin = section.right_margin = Cm(1.4)
        section.top_margin, section.bottom_margin = Cm(1.2), Cm(1.3)
    document.core_properties.author = 'XCalificator'
    document.core_properties.last_modified_by = 'XCalificator'
    title = _text(content.get('titulo') or material.get('titulo') or 'Evaluación')
    document.add_paragraph(title, 'Title')
    if soluciones or not print_settings:
        document.add_paragraph('Solucionario docente' if soluciones else 'Evaluación para estudiantes')
    if '_print_common_score' in content:
        document.add_paragraph(f'Cada pregunta vale {content["_print_common_score"]} puntos')
    document.add_paragraph('Nombre: __________________________________   Grado: __________')
    document.add_paragraph('Fecha: __________________   Nota: __________')
    for key, label in (('objetivo', 'Objetivo'), ('instrucciones', 'Instrucciones'), ('estrategia_lectora', 'Estrategia')):
        if content.get(key):
            document.add_paragraph(f'{label}: {_text(content[key])}')
    if content.get('texto'):
        document.add_paragraph('Texto de lectura', 'Heading 1')
        document.add_paragraph(_text(content['texto']))
        if content.get('fuente'):
            document.add_paragraph(f'Fuente: {_text(content["fuente"])}')
    for index, (event, raw, label) in enumerate(document_print_events(content, questions), 1):
        if event == 'seccion':
            subtotal = content.get('subtotales_secciones', {}).get(str(raw.id))
            title = ' '.join(value for value in (raw.etiqueta_original, raw.titulo) if value)
            heading = document.add_paragraph(title + (f' ({subtotal} puntos)' if subtotal is not None else ''),
                'Heading 2' if raw.padre_id else 'Heading 1')
            heading.paragraph_format.keep_with_next = True
            continue
        if event == 'texto':
            paragraph = document.add_paragraph(raw.texto)
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.keep_together = False
            continue
        if event == 'material':
            append_visual_docx(document, raw)
            continue
        question_start = len(document.element.body) - 1  # Inserted before sectPr.
        raw = raw if isinstance(raw, dict) else {'enunciado': str(raw)}
        safe = sanitize_student_payload(raw)
        number = label or raw.get('numero', index)
        score = raw.get('puntaje')
        suffix = f' ({score} puntos)' if score is not None and '_print_common_score' not in content else ''
        target, image_target = document, document
        available_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
        if raw.get('_print_lateral'):
            layout = next(iter(raw['_print_visual_layouts'].values()))
            width = available_cm * layout['ancho_porcentaje'] / 100
            frame = document.add_table(rows=1, cols=2)
            frame.autofit = False
            frame.columns[0].width, frame.columns[1].width = Cm(available_cm - width), Cm(width)
            target, image_target = frame.rows[0].cells
            target.width, image_target.width = Cm(available_cm - width), Cm(width)
            from docx.oxml import OxmlElement
            for cell in (target, image_target):
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'bottom', 'left', 'right'):
                    margin = OxmlElement('w:' + edge)
                    margin.set(qn('w:w'), '0')
                    margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                cell._tc.get_or_add_tcPr().append(margins)
        paragraph = target.paragraphs[0] if target is not document else document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(f'{number}. {_text(safe.get("enunciado"))}{suffix}').bold = True
        for visual in raw.get('_evaluation_visuals') or []:
            if not isinstance(visual, ResolvedVisualBlock):
                raise ValueError('El material visual necesita revisión antes de descargar.')
            layout = (raw.get('_print_visual_layouts') or {}).get(str(visual.block.id))
            append_visual_docx(image_target, visual,
                width_cm=max(1, available_cm * layout['ancho_porcentaje'] / 100 - (0.3 if raw.get('_print_lateral') else 0)) if layout else None,
                alignment=layout['alineacion'] if layout else None,
                include_caption=not raw.get('_print_lateral'))
            if raw.get('_print_lateral') and (visual.block.pie or visual.block.descripcion):
                caption = document.add_paragraph(visual.block.pie or visual.block.descripcion)
                caption.paragraph_format.space_after = Pt(2)
        options = [re.sub(r'^\s*[A-Za-z][).]\s*', '', _text(option)) for option in safe.get('opciones') or []]
        columns = raw.get('_print_option_columns', option_columns(options) if compact else 1)
        if columns == 2:
            table = target.add_table(rows=(len(options) + 1) // 2, cols=2)
            for index_option, option in enumerate(options):
                paragraph = table.cell(index_option // 2, index_option % 2).paragraphs[0]
                paragraph.text = f'{chr(65 + index_option)}) {option}'
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.keep_with_next = index_option < len(options) - 1
        else:
            for option_index, option in enumerate(options):
                p = target.add_paragraph(f'{chr(65 + option_index)}) {option}')
                p.paragraph_format.left_indent = Cm(0.5)
        if soluciones:
            answer = next((raw[k] for k in ('respuesta_correcta', 'respuesta_esperada', 'respuesta')
                           if raw.get(k) not in (None, '', [])), MISSING_ANSWER)
            target.add_paragraph(f'Respuesta: {_text(answer)}')
            for key, label in (('evidencia_textual', 'Evidencia'), ('justificacion', 'Por qué'), ('criterio_logro', 'Criterio')):
                if raw.get(key):
                    target.add_paragraph(f'{label}: {_text(raw[key])}')
        elif not safe.get('opciones'):
            try:
                lines = max(1, min(8, int(raw.get('lineas_respuesta') or (3 if kind in {'taller', 'lectura_comprensiva'} else 2))))
            except (TypeError, ValueError):
                lines = 3
            lines = raw.get('_print_lines', lines)
            for line in range(lines):
                paragraph = target.add_paragraph('' if print_settings else '________________________________________________________________')
                if print_settings:
                    from docx.oxml import OxmlElement
                    from docx.oxml.ns import qn
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = Pt(17 if compact else 23)
                    borders = OxmlElement('w:pBdr')
                    bottom = OxmlElement('w:bottom')
                    for key, value in [('val', 'single'), ('sz', '4'), ('color', '777777')]:
                        bottom.set(qn('w:' + key), value)
                    borders.append(bottom)
                    between = OxmlElement('w:between')
                    for key, value in [('val', 'single'), ('sz', '4'), ('color', '777777')]:
                        between.set(qn('w:' + key), value)
                    borders.append(between)
                    paragraph._p.get_or_add_pPr().append(borders)
                paragraph.paragraph_format.keep_with_next = line < lines - 1
            if raw.get('_print_drawing_mm'):
                from docx.enum.table import WD_ROW_HEIGHT_RULE
                box = (target.add_table(rows=1, cols=1, width=Cm(available_cm / 2))
                    if target is not document else target.add_table(rows=1, cols=1))
                box.style = 'Table Grid'
                box.rows[0].height = Cm(raw['_print_drawing_mm'] / 10)
                box.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if compact or print_settings:
            _keep_short_word_question(document, question_start, raw)
    if soluciones and content.get('criterios_revision'):
        document.add_paragraph('Criterios de revisión', 'Heading 1')
        for criterion in content['criterios_revision']:
            document.add_paragraph(_text(criterion))
    scores = [q.get('puntaje') if isinstance(q, dict) else None for q in questions]
    total = content.get('total_puntaje', content.get('puntaje_total'))
    if all(score is not None for score in scores):
        try:
            total = sum(Decimal(str(score)) for score in scores)
        except InvalidOperation:
            pass
    if total is not None:
        document.add_paragraph(f'Total: {total} puntos')
    output = BytesIO()
    document.save(output)
    return output.getvalue()
