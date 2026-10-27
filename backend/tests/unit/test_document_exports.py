from copy import deepcopy
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

from docx import Document
import pytest

from app.modules.evaluaciones.export_service import build_evaluation_material, render_evaluation_docx
from app.modules.herramientas.pdf_render import _render_examen
from app.modules.evaluaciones.print_layout import prepare_print_material
from tests.fixtures.document_structure import sample_document


def example(mode='compacta'):
    document, questions = sample_document()
    return SimpleNamespace(nombre='El jardín y sus semillas', descripcion='Lee y responde.',
        created_at=datetime(2026, 10, 26), nota_maxima=5, preguntas=questions,
        estructura_documento=document, lectura='LEGACY_DUPLICATE_SENTINEL',
        configuracion_impresion={'schema_version': 1, 'distribucion': mode},
        respuestas_esperadas=[{'numero': q['numero'], 'respuesta': f'PRIVATE_KEY_{q["numero"]}'} for q in questions])


@pytest.mark.parametrize('mode', ['compacta', 'media', 'amplia'])
def test_pdf_and_word_keep_whole_document_in_same_order_without_private_source(mode):
    item = example(mode)
    before = deepcopy(vars(item))
    material = build_evaluation_material(item, soluciones=False)
    html = _render_examen(prepare_print_material(material)['contenido_json'], False)
    doc = Document(BytesIO(render_evaluation_docx(material, soluciones=False)))
    word = '\n'.join(p.text for p in doc.paragraphs)
    for output in (html, word):
        assert output.index('Observamos el jardín') < output.index('Explica la observación científica 1.')
        assert output.index('Una semilla llegó al jardín') < output.index('Lo que dice el texto')
        assert output.index('Conversamos en grupo') < output.index('Mi conclusión')
        assert output.count('Una semilla llegó al jardín') == 1
        assert output.count('Explica la observación científica') == 16
        assert 'Comparen sus ideas' in output
        assert 'LEGACY_DUPLICATE_SENTINEL' not in output
        assert 'PRIVATE_KEY' not in output and 'source:' not in output
    assert vars(item) == before


def test_document_solutions_keep_canonical_keys_even_when_printed_labels_restart():
    material = build_evaluation_material(example(), soluciones=True)
    doc = Document(BytesIO(render_evaluation_docx(material, soluciones=True)))
    word = '\n'.join(p.text for p in doc.paragraphs)
    assert word.index('PRIVATE_KEY_14') > word.index('Mi conclusión')
    assert word.count('1. Explica') == 4


def test_unreviewed_or_incomplete_document_cannot_be_exported_as_complete():
    item = example()
    item.estructura_documento['revision'].update(estado='requiere_revision', confirmado=False)
    with pytest.raises(ValueError, match='[Rr]evis'):
        build_evaluation_material(item, soluciones=False)


def test_confirmed_section_totals_are_derived_from_canonical_question_maxima():
    from tests.unit.test_section_scoring import example as scoring_example
    from app.modules.evaluaciones.document_structure import compile_section_scores
    document, questions, _ = scoring_example()
    item = example()
    item.preguntas = compile_section_scores(document, questions, 5)
    raw = document.model_dump(mode='json'); raw['revision'].update(estado='completa', confirmado=True)
    item.estructura_documento = raw
    material = build_evaluation_material(item, soluciones=False)
    assert sorted(material['contenido_json']['subtotales_secciones'].values()) == ['1.00', '2.00', '2.00']
    html = _render_examen(prepare_print_material(material)['contenido_json'], False)
    word = '\n'.join(p.text for p in Document(BytesIO(render_evaluation_docx(material, soluciones=False))).paragraphs)
    assert 'Sección 1 (2.00 puntos)' in html and 'Sección 3 (1.00 puntos)' in word
