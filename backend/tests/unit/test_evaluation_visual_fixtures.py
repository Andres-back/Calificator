"""Smoke checks for the independently declared synthetic regression corpus."""
from io import BytesIO

import fitz
from docx import Document

from tests.fixtures.evaluation_visuals import (
    all_visual_documents,
    photo,
    shared_multipage_pdf,
    word_embedded_tables,
)


def test_six_synthetic_classes_and_inventories_are_distinct_and_readable():
    samples = all_visual_documents()
    assert len(samples) == 6
    assert len({sample.name for sample in samples}) == 6
    for sample in samples:
        assert sample.content and sample.elements and sample.question_numbers
        if sample.mime.endswith("document"):
            assert Document(BytesIO(sample.content)).paragraphs
        elif sample.mime == "application/pdf":
            with fitz.open(stream=sample.content, filetype="pdf") as document:
                assert document.page_count >= 1
        else:
            assert sample.content.startswith(b"\x89PNG")


def test_embedded_tables_are_images_not_native_word_tables_with_shared_context():
    sample = word_embedded_tables()
    document = Document(BytesIO(sample.content))
    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 5
    assert [element.kind for element in sample.elements] == ["tabla"] * 4 + ["grafica"]
    assert sample.elements[0].question_numbers == (1, 2, 3)
    assert sample.elements[-1].values == ("6", "3", "1")


def test_shared_source_stays_on_first_page_and_questions_on_second():
    sample = shared_multipage_pdf()
    with fitz.open(stream=sample.content, filetype="pdf") as document:
        assert document.page_count == 2
        assert document[0].get_images() and document[0].get_drawings()
        assert "1. Consulte" not in document[0].get_text()
        assert all(f"{number}. Consulte" in document[1].get_text() for number in (1, 2, 3))
    assert all(element.source_page == 1 for element in sample.elements)


def test_incorrect_annotations_are_explicit_variants_not_expected_solutions():
    sample = word_embedded_tables(annotations=True)
    text = "\n".join(p.text for p in Document(BytesIO(sample.content)).paragraphs)
    assert sample.has_student_annotations and "90% (incorrecta)" in text
    assert sample.elements[-1].values == ("6", "3", "1")
    assert photo(annotations=True).has_student_annotations
