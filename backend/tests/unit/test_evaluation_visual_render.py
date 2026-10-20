from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import UUID

from docx import Document
from PIL import Image
import pytest

from app.modules.evaluaciones import export_service, visual_render
from app.modules.evaluaciones.visual_schemas import VisualChart, VisualManifest, VisualTable
from tests.unit.test_evaluation_exports import evaluation
from tests.unit.test_evaluation_visual_contract import manifest_payload


def test_table_and_chart_raster_are_deterministic_and_keep_data():
    table = VisualTable(encabezados=['Producto', 'Precio'], filas=[['Libro', '2400'], ['Lápiz', '800']], unidad='pesos')
    chart = VisualChart(tipo='barras', categorias=['Libro', 'Lápiz'], series=[{'nombre': 'Precio', 'valores': [2400, 800]}], unidad='pesos')
    for value in (table, chart):
        first = visual_render.render_structured_visual(value)
        assert first == visual_render.render_structured_visual(value)
        with Image.open(BytesIO(first)) as image:
            assert image.format == 'PNG' and image.width >= 1000


@pytest.mark.asyncio
async def test_export_resolver_uses_revision_scoped_resources_and_fails_if_missing(monkeypatch):
    payload = manifest_payload()
    resource_id = UUID(payload['recurso_ids'][0])
    data = visual_render.render_structured_visual(VisualTable.model_validate(payload['bloques'][0]['tabla']))
    item = evaluation()
    revision = SimpleNamespace(snapshot_json={'contenido_visual': payload}, recurso_ids=[str(resource_id)])
    resource = SimpleNamespace(id=resource_id, evaluacion_id=item.id, estado='ready', clase='render_tabla', seguridad_contenido='publicable_confirmado', mime='image/png')
    db = SimpleNamespace(scalars=AsyncMock(return_value=[resource]))
    read = AsyncMock(return_value=data)
    monkeypatch.setattr(visual_render, 'read_resource_bytes', read)
    blocks = await visual_render.resolve_visual_material(db, item.id, revision)
    assert blocks[0].block.tabla.filas[0] == ['2021', '2400']
    read.assert_awaited_once_with(resource)
    db.scalars.return_value = []
    with pytest.raises(ValueError, match='material'):
        await visual_render.resolve_visual_material(db, item.id, revision)
    assert read.await_count == 1


def test_shared_table_is_once_in_word_and_html_and_images_embed_without_urls():
    payload = manifest_payload()
    manifest = VisualManifest.model_validate(payload)
    table = manifest.bloques[0]
    block = visual_render.ResolvedVisualBlock(table, visual_render.render_structured_visual(table.tabla))
    item = evaluation()
    for question, identifier in zip(item.preguntas, payload['pregunta_ids']):
        question['contenido_id'] = identifier
    material = export_service.build_evaluation_material(item, soluciones=False, visual_blocks=[block])
    document = Document(BytesIO(export_service.render_evaluation_docx(material, soluciones=False)))
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == '2400'
    assert document.tables[0].rows[0]._tr.xml.count('tblHeader') >= 1
    from app.modules.herramientas.pdf_render import render_material_html
    html = render_material_html(material, soluciones=False)
    assert html.count('>2400<') == 1
    assert 'body/p[2]' not in html and 'privado_original' not in html
    image = table.model_copy(update={'tipo': 'imagen', 'tabla': None})
    material = export_service.build_evaluation_material(item, soluciones=False,
        visual_blocks=[visual_render.ResolvedVisualBlock(image, block.image_bytes)])
    document = Document(BytesIO(export_service.render_evaluation_docx(material, soluciones=False)))
    assert len(document.inline_shapes) == 1
    assert 'data:image/png;base64,' in render_material_html(material, soluciones=False)


def test_incomplete_or_unconfirmed_visual_context_never_exports_as_complete():
    payload = manifest_payload()
    payload['bloques'][0]['seguridad_contenido'] = 'privado_original'
    block = VisualManifest.model_validate(payload).bloques[0]
    with pytest.raises(ValueError, match='(?i)revis'):
        visual_render.ResolvedVisualBlock(block, b'not-png')
    with pytest.raises(ValueError, match='(?i)divid'):
        visual_render.render_structured_visual(VisualTable(encabezados=['A'] * 50, filas=[['0'] * 50]))
