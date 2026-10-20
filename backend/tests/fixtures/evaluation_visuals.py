"""Synthetic documents and independent, hand-declared visual inventories.

No teacher documents, learner identities, answer keys or external downloads are
used. Constructors return bytes in memory; test callers own their temporary files.
The inventories are not inferred from the document or by an AI under test.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from math import cos, pi, sin


@dataclass(frozen=True)
class ExpectedElement:
    key: str
    kind: str
    question_numbers: tuple[int, ...]
    labels: tuple[str, ...]
    values: tuple[str, ...]
    source_page: int | None = None


@dataclass(frozen=True)
class VisualDocumentFixture:
    name: str
    mime: str
    content: bytes
    elements: tuple[ExpectedElement, ...]
    question_numbers: tuple[int, ...]
    has_student_annotations: bool = False


# Deliberately artificial values, different from the supplied teacher document.
TABLES = (
    (("Periodo", "Precio (pesos)"), ("2021", "2400"), ("2022", "2600")),
    (("Producto", "Cantidad", "Precio (pesos)"), ("Manzanas", "3", "1400")),
    (("Semana", "Asistentes"), ("1", "6"), ("2", "12"), ("3", "24")),
    (("Dia", "Hora", "Actividad"), ("Sabado", "09:00", "Lectura")),
)
INVENTORY = (
    ExpectedElement("precios", "tabla", (1, 2, 3), TABLES[0][0], ("2021", "2400", "2022", "2600")),
    ExpectedElement("compra", "tabla", (4,), TABLES[1][0], ("Manzanas", "3", "1400")),
    ExpectedElement("asistencia", "tabla", (5,), TABLES[2][0], ("1", "6", "2", "12", "3", "24")),
    ExpectedElement("horario", "tabla", (5,), TABLES[3][0], ("Sabado", "09:00", "Lectura")),
    ExpectedElement("distribucion", "grafica", (6,), ("Grupo A", "Grupo B", "Grupo C"), ("6", "3", "1")),
)


def _table_png(rows: tuple[tuple[str, ...], ...]) -> bytes:
    import fitz

    with fitz.open() as document:
        page = document.new_page(width=480, height=42 * len(rows) + 20)
        cell_width = 460 / len(rows[0])
        for row_index, row in enumerate(rows):
            for column_index, text in enumerate(row):
                x, y = 10 + column_index * cell_width, 10 + row_index * 42
                page.draw_rect(fitz.Rect(x, y, x + cell_width, y + 42), color=(0, 0, 0))
                page.insert_text((x + 7, y + 25), text, fontsize=11)
        return page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")


def _pie_png() -> bytes:
    import fitz

    with fitz.open() as document:
        page = document.new_page(width=480, height=230)
        start = 0.0
        for label, amount, color in (
            ("Grupo A: 6", 6, (0.2, 0.5, 0.9)),
            ("Grupo B: 3", 3, (0.3, 0.8, 0.5)),
            ("Grupo C: 1", 1, (0.9, 0.6, 0.2)),
        ):
            end = start + 2 * pi * amount / 10
            points = [fitz.Point(115, 115)]
            points.extend(fitz.Point(115 + 95 * cos(start + (end - start) * i / 60),
                                     115 + 95 * sin(start + (end - start) * i / 60))
                          for i in range(61))
            points.append(points[0])
            page.draw_polyline(points, color=(0, 0, 0), fill=color, closePath=True)
            page.insert_text((235, 70 + start * 15), label, fontsize=13)
            start = end
        return page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")


def _word(*, native_tables: bool, annotations: bool = False) -> bytes:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading("Evaluacion sintetica de datos", level=1)
    document.add_paragraph("Use la tabla de precios para las preguntas 1 a 3.")
    for index, rows in enumerate(TABLES):
        if native_tables:
            table = document.add_table(rows=0, cols=len(rows[0]))
            for values in rows:
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = value
        else:
            document.add_picture(BytesIO(_table_png(rows)), width=Inches(5))
        if index == 0:
            for number in (1, 2, 3):
                document.add_paragraph(f"{number}. Consulte los precios y explique su respuesta.")
        elif index in (1, 2):
            document.add_paragraph(f"{index + 3}. Interprete los datos de la tabla.")
    document.add_picture(BytesIO(_pie_png()), width=Inches(5))
    document.add_paragraph("6. Que porcentaje corresponde al Grupo A?")
    if annotations:
        document.add_paragraph("Anotacion ficticia del alumno: Grupo A = 90% (incorrecta).")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def word_native_tables() -> VisualDocumentFixture:
    return VisualDocumentFixture("word-tablas-nativas", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 _word(native_tables=True), INVENTORY, (1, 2, 3, 4, 5, 6))


def word_embedded_tables(*, annotations: bool = False) -> VisualDocumentFixture:
    return VisualDocumentFixture("word-cuatro-png-y-circular", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 _word(native_tables=False, annotations=annotations), INVENTORY,
                                 (1, 2, 3, 4, 5, 6), annotations)


def _mixed_pdf(*, separate_pages: bool = False) -> bytes:
    import fitz

    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((40, 40), "Use la tabla para responder las preguntas 1 a 3.")
        page.insert_image(fitz.Rect(40, 60, 520, 210), stream=_table_png(TABLES[0]))
        page.insert_text((40, 250), "Grafico vectorial: cajas (unidades)")
        for index, (label, value) in enumerate((("A", 4), ("B", 8))):
            x = 70 + index * 120
            page.draw_rect(fitz.Rect(x, 410 - value * 15, x + 70, 410), fill=(0.2, 0.5, 0.8))
            page.insert_text((x, 435), f"{label}: {value}")
        if separate_pages:
            page = document.new_page()
        for number in (1, 2, 3):
            y = (500 if not separate_pages else 80) + (number - 1) * 55
            page.insert_text((40, y), f"{number}. Consulte los precios en la tabla.")
            page.insert_text((40, y + 22), f"Espacio de respuesta {number}")
        return document.tobytes()


PDF_INVENTORY = (
    ExpectedElement("precios", "tabla", (1, 2, 3), TABLES[0][0], ("2021", "2400", "2022", "2600"), 1),
    ExpectedElement("cajas", "grafica_vectorial", (1, 2, 3), ("cajas", "unidades", "A", "B"), ("4", "8"), 1),
)


def mixed_pdf() -> VisualDocumentFixture:
    return VisualDocumentFixture("pdf-mixto-vectorial", "application/pdf", _mixed_pdf(), PDF_INVENTORY, (1, 2, 3))


def shared_multipage_pdf() -> VisualDocumentFixture:
    return VisualDocumentFixture("pdf-estimulo-otra-pagina", "application/pdf", _mixed_pdf(separate_pages=True), PDF_INVENTORY, (1, 2, 3))


def photo(*, annotations: bool = False) -> VisualDocumentFixture:
    import fitz

    with fitz.open(stream=_mixed_pdf(), filetype="pdf") as document:
        if annotations:
            document[0].insert_text((80, 600), "Respuesta ficticia incorrecta: 99", color=(0, 0, 1))
        content = document[0].get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")
    return VisualDocumentFixture("foto-hoja-sintetica", "image/png", content, PDF_INVENTORY, (1, 2, 3), annotations)


def scanned_pdf() -> VisualDocumentFixture:
    import fitz

    with fitz.open() as document:
        page = document.new_page()
        page.insert_image(page.rect, stream=photo().content)
        content = document.tobytes()
    return VisualDocumentFixture("pdf-escaneado", "application/pdf", content, PDF_INVENTORY, (1, 2, 3))


def all_visual_documents() -> tuple[VisualDocumentFixture, ...]:
    return (word_native_tables(), word_embedded_tables(), mixed_pdf(), scanned_pdf(), photo(), shared_multipage_pdf())
